import streamlit as st
import pandas as pd
import base64
from io import BytesIO
import openpyxl

# Set page config
st.set_page_config(page_title="Excel Editor", layout="wide")

# Custom CSS for styling
st.markdown("""
<style>
    .main .block-container {
        padding-top: 2rem;
    }
    h1 {
        color: #2c3e50;
        text-align: center;
        margin-bottom: 2rem;
    }
    .stButton > button {
        width: 100%;
    }
    .excel-table {
        border-collapse: collapse;
        width: 100%;
    }
    .excel-table th, .excel-table td {
        border: 1px solid #ddd;
        padding: 8px;
        text-align: left;
    }
    .excel-table tr:nth-child(even) {
        background-color: #f2f2f2;
    }
    .excel-table th {
        padding-top: 12px;
        padding-bottom: 12px;
        background-color: #4CAF50;
        color: white;
    }
</style>
""", unsafe_allow_html=True)

# Title
st.title("Interactive Excel Editor")

# Function to create HTML representation of Excel structure
def create_excel_structure_html(sheet):
    html = "<table class='excel-table'>"
    merged_cells = sheet.merged_cells.ranges

    for row in sheet.iter_rows():
        html += "<tr>"
        for cell in row:
            # Check if the cell is part of a merged range
            merged = False
            for merged_range in merged_cells:
                if cell.coordinate in merged_range:
                    if cell.coordinate == merged_range.start_cell.coordinate:
                        rowspan = merged_range.max_row - merged_range.min_row + 1
                        colspan = merged_range.max_col - merged_range.min_col + 1
                        html += f"<td rowspan='{rowspan}' colspan='{colspan}'>{cell.value}</td>"
                    merged = True
                    break
            if not merged:
                html += f"<td>{cell.value}</td>"
        html += "</tr>"
    html += "</table>"
    return html

# File uploader
uploaded_file = st.file_uploader("Choose an Excel file", type="xlsx")

if uploaded_file is not None:
    # Read Excel file
    excel_file = openpyxl.load_workbook(uploaded_file)
    sheet = excel_file.active

    # Display original Excel structure
    st.subheader("Original Excel Structure")
    excel_html = create_excel_structure_html(sheet)
    st.markdown(excel_html, unsafe_allow_html=True)

    # Read as pandas DataFrame for further processing
    df = pd.read_excel(uploaded_file)
    
    # Column selection for deletion
    st.subheader("Select columns to delete")
    cols_to_delete = st.multiselect("Choose columns to remove", df.columns.tolist())
    
    if cols_to_delete:
        df = df.drop(columns=cols_to_delete)
        st.success(f"Deleted columns: {', '.join(cols_to_delete)}")
    
    # Row deletion
    st.subheader("Delete rows")
    num_rows = st.number_input("Enter the number of rows to delete from the start", min_value=0, max_value=len(df)-1, value=0)
    
    if num_rows > 0:
        df = df.iloc[num_rows:]
        st.success(f"Deleted first {num_rows} rows")
    
    # Display editable dataframe
    st.subheader("Edit Data")
    st.write("You can edit individual cell values directly in the table below:")
    
    # Convert dataframe to a dictionary for st.data_editor
    data_dict = df.to_dict('list')
    edited_data = st.data_editor(data_dict)
    
    # Convert edited data back to dataframe
    edited_df = pd.DataFrame(edited_data)
    
    # Display edited data
    st.subheader("Edited Data")
    st.dataframe(edited_df)
    
    # Download button
    def get_excel_download_link(df):
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False)
        excel_data = output.getvalue()
        b64 = base64.b64encode(excel_data).decode()
        return f'<a href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64}" download="edited_file.xlsx">Download Edited Excel File</a>'
    
    st.markdown(get_excel_download_link(edited_df), unsafe_allow_html=True)

else:
    st.info("Please upload an Excel file to begin editing.")

# Add some final instructions
st.markdown("""
---
### Instructions:
1. Upload an Excel file using the file uploader at the top of the page.
2. View the original Excel structure, including merged cells.
3. Select columns to delete if needed.
4. Specify the number of rows to delete from the start, if any.
5. Edit individual cell values directly in the editable table.
6. Review your changes in the "Edited Data" section.
7. Download the edited Excel file using the download link provided.

Enjoy editing your Excel files with ease!
""")
